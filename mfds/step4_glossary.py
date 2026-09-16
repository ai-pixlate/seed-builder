"""4단계 — 식약처 전량을 용어집_성분 형식(매칭후보 탭과 같은 열·색)으로 만든다.

실행:  python step4_glossary.py
입력:  raw/mfds_raw.jsonl, 용어집_성분.xlsx (머리글 모양만 참고)
출력:  용어집_성분_2026-09-14.xlsx
"""
import copy
import json
import re
from collections import Counter
from pathlib import Path

import openpyxl
from openpyxl.styles import Font

HERE = Path(__file__).parent
RAW = HERE / "raw" / "mfds_raw.jsonl"
STYLE_SRC = HERE / "용어집_성분.xlsx"
OUT = HERE / "용어집_성분_2026-09-14.xlsx"

SOURCE = "식약처 화장품 원료성분정보 API(15111774)"
BASE_NOTE = "식약처 표준영문명"

# 한 칸 안의 여러 이름 구분: 쉼표. 단 양쪽이 모두 숫자인 쉼표(1,2-Hexanediol)는 이름의 일부.
NAME_SEP = re.compile(r"(?<!\d),|,(?!\d)")
SYN_SEP = re.compile(r"(?<!\d),|,(?!\d)|\r?\n")
CI_RE = re.compile(r"\bCI\s?\d{5}")

TAG_MULTI = "INCI 확인 필요: 영문명 여러 개"
TAG_CI = "INCI 확인 필요: 색소 CI 표기"
TAG_CAPS = "INCI 확인 필요: 대문자 표기"
TAG_NOEN = "INCI 확인 필요: 영문명 없음"

# 규제는 용어집 판단 대상이 아님 — 참고로만 note 에 남김 (enforcement 는 강제 유지)
REG_NOTES = {
    "황색4호": "규제 참고: 착색제 — 미국 라벨은 FD&C Yellow No. 5 주표기, CI 번호는 괄호 병기로 알려짐, 규제집 확인",
    "디에칠아미노하이드록시벤조일헥실벤조에이트": "규제 참고: 자외선차단 필터(DHHB) — 미국 자외선차단 용도 미승인으로 알려짐, 규제집 확인",
    "에칠헥실트리아존": "규제 참고: 자외선차단 필터 — 미국 모노그래프 미등재로 알려짐, 규제집 확인",
}


def split_names(s, sep=NAME_SEP):
    return [p.strip() for p in sep.split(s) if p.strip()] if s else []


def is_caps(part):
    """INCI 는 보통 단어 첫 글자만 대문자. 전부 대문자인 긴 이름을 잡는다.
    PEG-25 PABA, TEA-EDTA 처럼 4글자 이하 약어로만 된 이름은 제외."""
    words = re.findall(r"[A-Za-z]+", part)
    letters = "".join(words)
    return (len(letters) >= 6 and letters.isupper() and not CI_RE.fullmatch(part)
            and any(len(w) > 4 for w in words))


def tags_for(r):
    en = r["INGR_ENG_NAME"]
    if not en or not en.strip():
        return [TAG_NOEN]
    parts = split_names(en)
    tags = []
    if len(parts) >= 2:
        tags.append(TAG_MULTI)
    if CI_RE.search(en):
        tags.append(TAG_CI)
    if any(is_caps(p) for p in parts):
        tags.append(TAG_CAPS)
    return tags


def main():
    recs = [json.loads(l) for l in RAW.read_text(encoding="utf-8").splitlines() if l.strip()]
    src = openpyxl.load_workbook(STYLE_SRC)["매칭후보"]
    heads = list(src[1][:16])  # 구역 ~ version
    green = src["Q1"]
    names = [h.value for h in heads]

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "용어집_성분"
    for j, h in enumerate(heads + [green], start=1):
        c = ws.cell(1, j, h.value if j <= 16 else "이명")
        c.font, c.fill = copy.copy(h.font), copy.copy(h.fill)
        c.alignment, c.border = copy.copy(h.alignment), copy.copy(h.border)
        letter = c.column_letter
        ws.column_dimensions[letter].width = src.column_dimensions[h.column_letter].width if j <= 16 else 40
    ws.column_dimensions["O"].width = 60

    tag_count = Counter()
    reg_hit = []
    for i, r in enumerate(recs, start=2):
        tags = tags_for(r)
        tag_count.update(tags)
        if not tags:
            tag_count["확인 표시 없음"] += 1
        note_parts = [BASE_NOTE] + tags
        if r["INGR_KOR_NAME"] in REG_NOTES:
            note_parts.append(REG_NOTES[r["INGR_KOR_NAME"]])
            reg_hit.append(r["INGR_KOR_NAME"])
        row = dict.fromkeys(names)
        row.update(source_ko=r["INGR_KOR_NAME"],
                   target_text=r["INGR_ENG_NAME"] if r["INGR_ENG_NAME"] and r["INGR_ENG_NAME"].strip() else None,
                   term_kind="ingredient", enforcement="강제", target_lang="en",
                   internal_category="common", source=SOURCE, note=" / ".join(note_parts), version=1)
        for j, n in enumerate(names, start=1):
            ws.cell(i, j, row[n])
        ws.cell(i, 17, r["INGR_SYNONYM"])
    ws.freeze_panes = "D2"
    ws.auto_filter.ref = f"A1:Q{len(recs) + 1}"

    n_en = sum(1 for r in recs if r["INGR_ENG_NAME"] and r["INGR_ENG_NAME"].strip())
    n_syn = sum(1 for r in recs if r["INGR_SYNONYM"])
    n_syn_multi = sum(1 for r in recs if len(split_names(r["INGR_SYNONYM"], SYN_SEP)) >= 2)
    n_check = len(recs) - tag_count["확인 표시 없음"]
    semicolon = [r["INGR_KOR_NAME"] for r in recs
                 if any(";" in (r[f] or "") for f in ("INGR_KOR_NAME", "INGR_ENG_NAME"))]

    rd = wb.create_sheet("00_읽는법", 0)
    B = None
    lines = [
        ("용어집 — 성분(ingredient) 식약처 기준", B),
        ("작성", "2026-09-14"),
        ("상태", "검수 전. verified_at 이 빈 행은 시드에 넣지 않는다"),
        ("원천", f"{SOURCE} 전량 {len(recs)}건 (2026-09-14 수집). 가공 없음"),
        ("형식", "기존 용어집_성분.xlsx 「매칭후보」 탭과 같은 열 순서·색"),
        ("범위", "한글 성분명 → 공식 영문명(INCI) 번역이 올바른지만 다룬다. 미국 등에서 쓸 수 있는 성분인지(규제)는 규제집에서 판단 — 여기서는 note 에 참고로만 표시"),
        ("■ 색", B),
        ("머리글 남색", "DB로 가는 열"),
        ("머리글 회색", "시트 전용 (로더 무시)"),
        ("머리글 초록", "검토용 (DB로 안 감). 시드 전달 시 삭제"),
        ("■ 열", B),
        ("구역", "비움 — 식약처 전량이라 A/B/C 구분 없음"),
        ("id", "비움 — GL 번호는 검수 확정 후 부여. 번호 재사용 금지"),
        ("source_ko", "식약처 표준 성분명 — API 원본 그대로 (옛 표기 에칠·메칠·디 등 포함)"),
        ("target_text", "식약처 영문명 — API 원본 그대로. API에 없으면 비움 (추측·번역으로 채우지 않음)"),
        ("term_kind", "ingredient"),
        ("enforcement", "강제 (전 행)"),
        ("target_lang", "en"),
        ("internal_category", "common"),
        ("example_sentence / frequency / corpus_size / corpus_version", "비움"),
        ("verified_at", "비움 — 검수자가 채움"),
        ("source", SOURCE),
        ("note", "꼬리표를 ' / ' 로 이어 붙임 (세미콜론은 시드 구분 기호라 쓰지 않음). 아래 ■ note 꼬리표 참고"),
        ("version", "1"),
        ("이명 (초록)", "식약처 API의 '다른 이름'(INGR_SYNONYM) 칸 — API 원본 그대로"),
        ("■ note 꼬리표 — 엑셀 거르기로 'INCI 확인 필요' 만 골라 검수", B),
        (BASE_NOTE, f"모든 행 기본 ({len(recs)}건)"),
        (TAG_MULTI, f"영문명 칸에 이름이 쉼표로 여러 개 → 대표 표기 선택 필요. 예: 정제수 Water,Aqua ({tag_count[TAG_MULTI]}건)"),
        (TAG_CI, f"색소 등 CI 번호가 붙은 것. INCI 에서 색소는 보통 CI 번호가 공식 이름 → 표기 확인. 예: 황색4호 Tartrazine,CI 19140 ({tag_count[TAG_CI]}건)"),
        (TAG_CAPS, f"전부 대문자인 이름이 섞인 것 (INCI 는 보통 첫 글자만 대문자). PEG·TEA 같은 짧은 약어만 대문자인 것은 제외. 예: 마그네슘가루 MAGNESIUM POWDER ({tag_count[TAG_CAPS]}건)"),
        (TAG_NOEN, f"식약처에 영문명이 없음 — INCI 여부 판단 불가. 예: 가공소금 ({tag_count[TAG_NOEN]}건)"),
        ("규제 참고: …", f"규제 관련 참고 메모 ({len(reg_hit)}건). enforcement 에 영향 없음 — 판단은 규제집"),
        ("합계", f"INCI 확인 필요 꼬리표가 하나 이상 붙은 행 {n_check}건 (한 행에 여러 꼬리표 가능)"),
        ("■ 규제 참고 행", B),
    ] + [(k, v) for k, v in REG_NOTES.items()] + [
        ("■ 이명이란", B),
        ("뜻", "같은 성분을 부르는 다른 이름. 표준명이 바뀌기 전 옛 이름이 들어 있는 경우가 있음 (전부 옛 이름인지는 확인 안 함)"),
        ("예", "source_ko 가지열매추출물 → 이명 가지추출물"),
        ("건수", f"{n_syn}개 성분에 이명 있음 (나머지는 빈칸). 그중 {n_syn_multi}개는 한 칸에 이름이 여러 개"),
        ("여러 개일 때", "쉼표로 구분 (예: 리치열매추출물,여지열매추출물). 단 1,2-헥산디올처럼 양쪽이 숫자인 쉼표는 이름의 일부. '/'도 이름의 일부(말로우꽃/잎/줄기추출물)"),
        ("쓰는 법", "올리브영·OCR 표기가 source_ko 에서 안 찾아지면 이명 칸에서 찾아본다. 찾으면 그 행의 source_ko(표준명)·target_text 를 쓴다"),
        ("주의", "이명은 source_ko 로 옮기지 않았다 — 용어집 행은 전부 식약처 표준명 기준. 이명 칸은 시드 전달 시 삭제"),
        ("■ 그 밖의 주의", B),
        ("옛 표기", "식약처는 에칠·메칠·디 등 옛 표기가 섞여 있어 에틸·메틸·다이 표기로 검색하면 안 찾아질 수 있음. 용어집에 없다 ≠ 없는 성분"),
        ("중복·충돌", "나일론6/12 와 나일론-6/12 두 줄 / 같은 영문명에 한글명 여러 개 28건(예: Panax Ginseng Root ← 인삼, 산양삼뿌리) / 같은 이명이 다른 성분 2개에 걸린 경우 66건"),
        ("줄바꿈·공백 든 값", "칸 안 줄바꿈 5건(해바라기씨왁스·그린와틀꽃왁스 영문명 등), 앞뒤 공백 2건(아세틸트라이펩타이드-146, 다이라우라미도아이소프로판올). 원본 그대로 둠 — 시드 전달 전 확인"),
        ("세미콜론", f"성분명·영문명 칸에 세미콜론 든 행 {len(semicolon)}건"),
    ]
    for k, v in lines:
        rd.append([k, v])
        if v is None:
            rd.cell(rd.max_row, 1).font = Font(bold=True)
    rd.column_dimensions["A"].width = 40
    rd.column_dimensions["B"].width = 110
    wb.save(OUT)

    print(f"행 {len(recs)} · 영문명 있음 {n_en} / 없음 {len(recs) - n_en}")
    for k in (TAG_MULTI, TAG_CI, TAG_CAPS, TAG_NOEN, "확인 표시 없음"):
        print(f"  {k}: {tag_count[k]}")
    print(f"  INCI 확인 필요 1개 이상: {n_check}")
    print(f"  규제 참고: {reg_hit}")
    print(f"  이명 있음 {n_syn} / 여러 개 {n_syn_multi} · 세미콜론 {semicolon}")


if __name__ == "__main__":
    main()
