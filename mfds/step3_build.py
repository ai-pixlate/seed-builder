"""3~6단계 — 매칭키 만들기, 대조본·식약처_사전 시트, 보고서, 검증.

실행:  python step3_build.py
입력:  raw/mfds_raw.jsonl, raw/checkpoint.json, 용어집_성분.xlsx (읽기만)
출력:  용어집_성분_대조본.xlsx, report_*.csv, run_log.md
"""
import copy
import csv
import json
import re
import time
import unicodedata
from collections import Counter, defaultdict
from pathlib import Path

import openpyxl
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE

HERE = Path(__file__).parent
RAW = HERE / "raw" / "mfds_raw.jsonl"
CKPT = HERE / "raw" / "checkpoint.json"
SRC_XLSX = HERE / "용어집_성분.xlsx"
OUT_XLSX = HERE / "용어집_성분_대조본.xlsx"

F_KO, F_EN, F_CAS, F_ORIGIN, F_SYN = (
    "INGR_KOR_NAME", "INGR_ENG_NAME", "CAS_NO", "ORIGIN_MAJOR_KOR_NAME", "INGR_SYNONYM")

# 이명 칸 안에서 여러 이름을 가르는 기호 — 전량 확인 결과 쉼표(161건)·줄바꿈(1건).
# "/" 는 성분명 자체에 들어가므로(말로우꽃/잎/줄기추출물) 가르지 않음.
# 양쪽이 모두 숫자인 쉼표(1,2-헥산다이올)만 이름의 일부라 가르지 않음 (CI 77480,Gold 는 가름).
SYN_SEP = r"(?<!\d),|,(?!\d)|\r?\n"

SOURCE = "식약처 화장품 원료성분정보 API(15111774)"
NOTE = "식약처 표준영문명 — 미국 라벨 대조 안 됨"

# 사용자 지정 플래그 — 식약처_사전 탭에서 note 에 사유, enforcement 비움
FLAGS = {
    "황색4호": "미국 라벨은 FD&C Yellow No. 5 주표기, CI 번호는 괄호 병기만 허용 — 식약처 영문명으로 대체 불가",
    "디에틸아미노하이드록시벤조일헥실벤조에이트": "자외선차단 필터(DHHB) — 미국에서 자외선차단 용도 미승인, 번역이 아니라 판매 가능성 검토 대상",
    "에틸헥실트리아존": "자외선차단 필터 — 미국 모노그래프 미등재로 알려짐, 확인 필요",
}
# 식약처 사전의 실제 표기 (옛 표기 칠/디 사용 — 전량 검색으로 1건씩 확인)
FLAG_MFDS_NAMES = {
    "디에틸아미노하이드록시벤조일헥실벤조에이트": "디에칠아미노하이드록시벤조일헥실벤조에이트",
    "에틸헥실트리아존": "에칠헥실트리아존",
}
# 식약처 영문명 칸 안에서 여러 영문명을 가르는 쉼표 (1,2-Hexanediol 같은 숫자 사이 쉼표는 제외)
EN_SEP = r"(?<!\d),|,(?!\d)"

# ---------------------------------------------------------------- 3단계: 매칭키

_FULLWIDTH = {cp: cp - 0xFEE0 for cp in range(0xFF01, 0xFF5F)}
_FULLWIDTH[0x3000] = 0x20
_DROP = re.compile(r"[\s\-·.]")


def ko_key(s):
    if s is None:
        return None
    s = unicodedata.normalize("NFC", str(s))
    s = s.translate(_FULLWIDTH)
    s = _DROP.sub("", s)
    return s.lower() or None


def split_syn(s):
    if not s:
        return []
    parts = [s] if SYN_SEP is None else re.split(SYN_SEP, s)
    return [p.strip() for p in parts if p and p.strip()]


# ---------------------------------------------------------------- 공통

def blank(v):
    return v is None or (isinstance(v, str) and v.strip() == "")


def write_csv(name, header, rows):
    with (HERE / name).open("w", encoding="utf-8-sig", newline="") as f:
        w = csv.writer(f)
        w.writerow(header)
        w.writerows(rows)


def split_en(s):
    return [p.strip() for p in re.split(EN_SEP, s) if p.strip()]


def en_diff_kind(a, alts):
    """기존 target_text(a) 와 식약처 영문명 후보들(alts) 중 가장 가까운 차이."""
    kinds = ["대소문자·띄어쓰기·기호만 다름", "괄호 병기만 다름", "다른 표기 (다른 성분일 수 있음)"]
    return min((_diff(a, b) for b in alts), key=kinds.index)


def _diff(a, b):
    def loose(x):
        return re.sub(r"[^0-9a-z]", "", x.lower())

    def no_paren(x):
        return re.sub(r"\s*\([^)]*\)", "", x).strip()

    if loose(a) == loose(b):
        return "대소문자·띄어쓰기·기호만 다름"
    if loose(no_paren(a)) == loose(no_paren(b)):
        return "괄호 병기만 다름"
    return "다른 표기 (다른 성분일 수 있음)"


def main():
    t0 = time.time()
    ckpt = json.loads(CKPT.read_text(encoding="utf-8"))
    recs = [json.loads(l) for l in RAW.read_text(encoding="utf-8").splitlines() if l.strip()]
    log = []

    # ---------- 색인
    std_idx = defaultdict(list)
    syn_idx = defaultdict(list)
    for r in recs:
        std_idx[ko_key(r[F_KO])].append(r)
        for s in split_syn(r[F_SYN]):
            syn_idx[ko_key(s)].append(r)

    # ---------- 4단계: 매칭후보 대조
    wb = openpyxl.load_workbook(SRC_XLSX)
    ws = wb["매칭후보"]
    head = [c.value for c in ws[1]]
    col = {h: i for i, h in enumerate(head)}
    green = ws.cell(1, col["올리브영_등장상품수"] + 1)
    new_cols = ["식약처_영문명", "식약처_일치여부", "식약처_표준한글명", "식약처_찾은방법"]
    start = ws.max_column + 1
    for j, name in enumerate(new_cols):
        c = ws.cell(1, start + j, name)
        c.font, c.fill, c.alignment, c.border = (copy.copy(green.font), copy.copy(green.fill),
                                                 copy.copy(green.alignment), copy.copy(green.border))
        ws.column_dimensions[openpyxl.utils.get_column_letter(start + j)].width = 28 if j != 1 else 16

    match_rows = []
    for i in range(2, ws.max_row + 1):
        zone = ws.cell(i, col["구역"] + 1).value
        gid = ws.cell(i, col["id"] + 1).value
        ko = ws.cell(i, col["source_ko"] + 1).value
        tt = ws.cell(i, col["target_text"] + 1).value
        k = ko_key(ko)
        hits, via = std_idx.get(k), "표준명"
        if not hits:
            hits, via = syn_idx.get(k), "이명"
        if not hits:
            en_join = std_join = None
            status, via = "식약처에없음", None
            ens = []
        else:
            ens = list(dict.fromkeys(h[F_EN] for h in hits if not blank(h[F_EN])))
            en_join = " | ".join(ens) or None
            std_join = " | ".join(dict.fromkeys(h[F_KO] for h in hits))
            if not ens:
                status = "식약처영문없음"
            elif blank(tt):
                status = "기존영문없음"
            elif tt in [a for e in ens for a in split_en(e)]:
                status = "일치"
            else:
                status = "불일치"
        for j, v in enumerate([en_join, status, std_join, via]):
            ws.cell(i, start + j, v)
        match_rows.append(dict(zone=zone, id=gid, ko=ko, tt=tt, en=en_join, ens=ens,
                               status=status, std=std_join, via=via))

    # ---------- 5단계: 식약처_사전 시트
    sd = wb.create_sheet("식약처_사전")
    base_cols = head[: col["version"] + 1]  # 구역 ~ version (매칭후보와 같은 순서)
    extra = ["cas_no", "기원및정의", "이명", "ko_key"]
    for j, name in enumerate(base_cols + extra, start=1):
        src = ws.cell(1, j) if j <= len(base_cols) else green
        c = sd.cell(1, j, name)
        c.font, c.fill, c.alignment, c.border = (copy.copy(src.font), copy.copy(src.fill),
                                                 copy.copy(src.alignment), copy.copy(src.border))
    for j in range(1, len(base_cols) + len(extra) + 1):
        letter = openpyxl.utils.get_column_letter(j)
        sd.column_dimensions[letter].width = ws.column_dimensions[letter].width or 14
    sd.freeze_panes = "C2"

    flag_hit = {name: [] for name in FLAGS}
    illegal = []
    for n, r in enumerate(recs, start=2):
        row = {h: None for h in base_cols}
        row.update(source_ko=r[F_KO], target_text=None if blank(r[F_EN]) else r[F_EN],
                   term_kind="ingredient", enforcement="강제", target_lang="en",
                   internal_category="common", source=SOURCE, note=NOTE, version=1)
        rk = ko_key(r[F_KO])
        syn_keys = {ko_key(s) for s in split_syn(r[F_SYN])}
        for name, reason in FLAGS.items():
            if (ko_key(name) == rk or ko_key(name) in syn_keys
                    or ko_key(FLAG_MFDS_NAMES.get(name)) == rk):
                row["enforcement"] = None
                row["note"] = f"플래그: {reason} / {NOTE}"
                flag_hit[name].append(r[F_KO])
        values = [row[h] for h in base_cols] + [r[F_CAS], r[F_ORIGIN], r[F_SYN], rk]
        for j, v in enumerate(values, start=1):
            if isinstance(v, str) and ILLEGAL_CHARACTERS_RE.search(v):
                illegal.append((n, (base_cols + extra)[j - 1]))
                raise SystemExit(f"엑셀에 넣을 수 없는 제어문자 발견: {r[F_KO]} / {(base_cols + extra)[j - 1]}"
                                 " — 원본을 임의로 고치지 않기 위해 멈춥니다.")
            sd.cell(n, j, v)

    wb.save(OUT_XLSX)

    # ---------- 보고서
    order = {"A": 0, "B": 1, "C": 2}
    conflicts = [m for m in match_rows if m["status"] == "불일치"]
    conflicts.sort(key=lambda m: (order.get(m["zone"], 9), m["id"] or ""))
    write_csv("report_conflict.csv",
              ["구역", "id", "source_ko", "기존_target_text", "식약처_영문명", "식약처_표준한글명",
               "찾은방법", "차이종류"],
              [[m["zone"], m["id"], m["ko"], m["tt"], m["en"], m["std"], m["via"],
                en_diff_kind(m["tt"], [a for e in m["ens"] for a in split_en(e)])]
               for m in conflicts])

    no_en = [r for r in recs if blank(r[F_EN])]
    write_csv("report_no_english.csv", ["식약처_표준한글명", "cas_no", "이명", "기원및정의"],
              [[r[F_KO], r[F_CAS], r[F_SYN], r[F_ORIGIN]] for r in no_en])

    coll = []
    for idx, kind in ((std_idx, "표준명"), (syn_idx, "이명")):
        for k, rs in idx.items():
            ens = list(dict.fromkeys(r[F_EN] for r in rs if not blank(r[F_EN])))
            if len(ens) >= 2:
                coll.append([kind, k, " | ".join(dict.fromkeys(r[F_KO] for r in rs)), " | ".join(ens), len(ens)])
    write_csv("report_collisions.csv", ["겹친곳", "ko_key", "식약처_표준한글명들", "영문명들", "영문명수"], coll)

    delim = []
    fields = [F_KO, F_EN, F_CAS, F_ORIGIN, F_SYN]
    for r in recs:
        for f in fields:
            v = r.get(f)
            if not isinstance(v, str):
                continue
            probs = [name for ch, name in ((";", "세미콜론"), ("\n", "줄바꿈"), ("\r", "줄바꿈(CR)"), ("\t", "탭"))
                     if ch in v]
            if probs:
                delim.append([r[F_KO], f, "+".join(probs), v])
    write_csv("report_delimiter.csv", ["식약처_표준한글명", "칸", "문제", "값"], delim)

    # ---------- 6단계: 검증 요약
    st = Counter((m["zone"], m["status"]) for m in match_rows)
    bc = [m for m in match_rows if m["zone"] in ("B", "C")]
    bc_found = [m for m in bc if m["status"] != "식약처에없음"]
    a_conf = [m for m in conflicts if m["zone"] == "A"]
    dup_names = {k: v for k, v in Counter(r[F_KO] for r in recs).items() if v > 1}
    delim_by = Counter((d[1], d[2]) for d in delim)
    delim_seed = [d for d in delim if d[1] in (F_KO, F_EN)]

    L = log.append
    L("# 식약처 성분 수집·대조 실행 기록\n")
    L(f"- 실행: {time.strftime('%Y-%m-%d %H:%M:%S')}")
    L(f"- API: `{SOURCE}` · 500건씩 {ckpt['last_page']}쪽")
    L(f"- 수집 시작 {ckpt.get('started_at')} → 끝 {ckpt.get('finished_at')} "
      f"(마지막 실행 {ckpt.get('elapsed_sec_this_run')}초) · 대조 단계 {round(time.time() - t0, 1)}초\n")
    L("## 1. 수집 건수")
    L(f"- totalCount **{ckpt['totalCount']}** / raw 줄 수 **{len(recs)}** → "
      f"{'일치' if ckpt['totalCount'] == len(recs) else '**불일치**'}")
    if ckpt.get("totalCount_changes"):
        L(f"- 수집 중 totalCount 변동: {ckpt['totalCount_changes']}")
    L(f"- 같은 표준한글명이 2번 이상 나온 것: {len(dup_names)}개 {list(dup_names)[:10]}")
    L(f"- 영문명 있음 {len(recs) - len(no_en)} / 없음 {len(no_en)} (→ report_no_english.csv)\n")
    L("## 2. B·C구역 42건 — 식약처에서 찾은 것")
    L(f"- **{len(bc_found)} / {len(bc)}건** (표준명 {sum(m['via'] == '표준명' for m in bc_found)}, "
      f"이명 {sum(m['via'] == '이명' for m in bc_found)})")
    for m in bc_found:
        L(f"  - [{m['zone']}] {m['ko']} → {m['en'] or '(식약처 영문명 없음)'} "
          f"· {m['status']} · {m['via']}{'(표준명: ' + m['std'] + ')' if m['via'] == '이명' else ''}")
    L("- 못 찾은 것: " + ", ".join(f"[{m['zone']}] {m['ko']}" for m in bc if m["status"] == "식약처에없음") + "\n")
    L("## 3. A구역 226건")
    L("- " + " · ".join(f"{s} {st[('A', s)]}" for s in
                        ("일치", "불일치", "식약처에없음", "식약처영문없음", "기존영문없음")))
    for m in a_conf:
        L(f"  - {m['id']} {m['ko']}: 기존 `{m['tt']}` ↔ 식약처 `{m['en']}` "
          f"— {en_diff_kind(m['tt'], [a for e in m['ens'] for a in split_en(e)])}")
    L("")
    L("## 4. 1:N 충돌 (report_collisions.csv)")
    L(f"- 표준명 기준 {sum(c[0] == '표준명' for c in coll)}건 · 이명 기준 {sum(c[0] == '이명' for c in coll)}건\n")
    L("## 5. 구분자·줄바꿈 (report_delimiter.csv)")
    L(f"- 값 안에 세미콜론/줄바꿈/탭이 든 칸: {len(delim)}건 — " +
      ", ".join(f"{f} {p} {n}" for (f, p), n in delim_by.items()))
    L(f"- 이 중 시드로 가는 source_ko·target_text 칸: **{len(delim_seed)}건**")
    for d in delim_seed[:30]:
        L(f"  - {d[0]!r} · {d[1]} · {d[2]} · `{d[3]!r}`")
    L("")
    L("## 6. 사용자 지정 플래그 (식약처_사전 탭 — note 사유, enforcement 비움)")
    for name, hit in flag_hit.items():
        L(f"- {name}: {'→ ' + ', '.join(hit) if hit else '식약처 사전에서 못 찾음 (적용 안 됨)'}")
    L("\n## 대조본 초록 열 값 뜻")
    L("- 식약처_일치여부: 일치 / 불일치 / 식약처에없음 / 식약처영문없음(식약처에 이름은 있으나 영문명 빈칸) "
      "/ 기존영문없음(기존 target_text 빈칸이라 비교 불가)")
    L("- 일치 판정: 식약처 영문명 칸에 쉼표로 여러 이름이 있으면(예: Water,Aqua) 그중 하나와 똑같으면 일치")
    L("- 식약처_찾은방법: 표준명 / 이명 (이명 = 식약처 INGR_SYNONYM 칸으로 찾음)")
    (HERE / "run_log.md").write_text("\n".join(log) + "\n", encoding="utf-8")
    print("\n".join(log))


if __name__ == "__main__":
    main()
